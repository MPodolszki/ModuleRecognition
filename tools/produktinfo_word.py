#!/usr/bin/env python3
"""The product texts as a Word document for marketing -- and back.

    uv run tools/produktinfo_word.py export      # -> dataset/products/Produktinfos.docx
    uv run tools/produktinfo_word.py import Produktinfos.docx

export: a title page (how to edit, the layout with the numbered fields),
then one page per product: a table Nr | Feld | Inhalt with the
current produktinfo.txt and a preview of its info panel underneath.
Marketing edits only the "Inhalt" column.

import: reads the tables back (one per product, recognised by its
"Artikelnummer" row) and rewrites each produktinfo.txt; prints what
changed. Afterwards: tools/render_panel.py for the new PDF,
tools/product_info.py + deploy.sh for the board.
"""

import argparse
import datetime
import io
import pathlib
import sys
import tempfile

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

from product_info import (FIELDS, INFO_FILE, SLOTS, build, product_list,  # noqa: E402
                          read_info_file, slot_keys, write_info_file)

PRODUCTS = ROOT / "dataset" / "products"
PETROL = "006E73"


def rows_for(row):
    """(nr, field, value, hint) per table row, one row per Vergleich /
    Alleinstellungsmerkmal place."""
    out = []
    for nr, field, hint in FIELDS:
        if field == "Artikelnummer":
            continue
        if field in SLOTS:
            for i in range(1, SLOTS[field] + 1):
                out.append((nr, f"{field} {i}", row.get(f"{field} {i}", ""),
                            hint if i == 1 else ""))
        elif field == "Neu":
            out.append((nr, "Neu", row.get("Neu (ja/nein)", ""), hint))
        else:
            out.append((nr, field, row.get(field, ""), hint))
    return out


def _shade(cell, color):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


def _cell_text(cell, text, size=9, bold=False, color=None):
    from docx.shared import Pt, RGBColor
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def export(out):
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401  (portrait is the default)
    from docx.shared import Cm, Pt, RGBColor

    import render_panel
    arts = product_list(PRODUCTS)
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading("Info-Fenster der Messe-Demo – Produkttexte", 0)
    doc.add_paragraph(f"Stand {datetime.date.today():%d.%m.%Y}. Pro Produkt eine Seite. "
                      "Bitte nur die Spalte „Inhalt“ ändern – Spalten „Nr“ und „Feld“ so "
                      "lassen, sie werden zum Einlesen gebraucht. Die Nummer zeigt, wo das "
                      "Feld im Fenster steht (Bild unten).")
    for tip in ("Leere Felder fallen im Fenster einfach weg.",
                "Kurzbeschreibung: höchstens 3 Zeilen im Fenster (ca. 270 Zeichen).",
                "Highlights: je eine kurze Zeile (ca. 55 Zeichen). „Vergleich 1–4“ "
                "dienen dem Vergleich zweier Produkte und stehen dann Zeile für Zeile "
                "nebeneinander – bitte bei allen Produkten dasselbe Thema auf denselben "
                "Platz (z. B. 1 CPU, 2 KI/NPU, 3 Speicher, 4 Schnittstellen). "
                "„Alleinstellungsmerkmal 1–4“: was nur dieses Produkt kann.",
                "Kerne, Größe, Temperatur, Verfügbar bis: kurz halten (ca. 16 Zeichen).",
                "Links: vollständige Adresse mit https://, daraus wird der QR-Code.",
                "Bemerkung wird nicht angezeigt – für Rückfragen und Quellen.",
                "Die Vorschau unter jeder Tabelle zeigt den Stand beim Erstellen dieses "
                "Dokuments, nicht die Änderungen. Nach dem Zurückschicken gibt es ein "
                "neues PDF mit allen Fenstern."):
        doc.add_paragraph(tip, style="List Bullet")
    doc.add_paragraph("Wo welches Feld im Fenster erscheint (Nummern wie in der Spalte „Nr“):")

    with tempfile.TemporaryDirectory() as tmp:
        infos = build(PRODUCTS, arts, pathlib.Path(tmp))
        for info in infos.values():
            if info.get("chart"):
                info["chart"] = str(pathlib.Path(tmp) / info["chart"])
        example = "PCM-070-0F3243I.A3" if "PCM-070-0F3243I.A3" in infos else arts[0]
        panel, marks = render_panel.render(example, infos[example])
        layout = render_panel.layout_page(panel, marks)
        buf = io.BytesIO()
        layout.save(buf, "PNG")
        buf.seek(0)
        doc.add_picture(buf, width=Cm(17.4))
        hints = doc.add_table(rows=0, cols=3)
        hints.style = "Table Grid"
        hints.autofit = False
        for column, w in zip(hints.columns, (Cm(0.9), Cm(4.4), Cm(12.1))):
            column.width = w
        for nr, field, hint in FIELDS:
            r = hints.add_row().cells
            for cell, w in zip(r, (Cm(0.9), Cm(4.4), Cm(12.1))):
                cell.width = w
            _cell_text(r[0], str(nr or "–"), 8, color="707070")
            _cell_text(r[1], field, 8, True)
            _cell_text(r[2], hint.replace(" -- ", " – "), 8)

        for art in arts:
            doc.add_page_break()
            path = PRODUCTS / art / INFO_FILE
            row = read_info_file(path) if path.exists() else {}
            doc.add_heading(row.get("Bezeichnung") or art, 1)
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            table.autofit = False  # fixed column widths (Word and LibreOffice)
            widths = (Cm(0.9), Cm(4.4), Cm(12.1))
            for column, w in zip(table.columns, widths):
                column.width = w
            for cell, text, w in zip(table.rows[0].cells, ("Nr", "Feld", "Inhalt"), widths):
                _cell_text(cell, text, 9, True, "FFFFFF")
                _shade(cell, PETROL)
                cell.width = w
            r = table.add_row().cells
            for cell, text, w in zip(r, ("4", "Artikelnummer", art), widths):
                _cell_text(cell, text, 9, color="707070")
                cell.width = w
            for nr, field, value, hint in rows_for(row):
                r = table.add_row().cells
                for cell, w in zip(r, widths):
                    cell.width = w
                _cell_text(r[0], str(nr or "–"), 9, color="707070")
                _cell_text(r[1], field, 9, True)
                _cell_text(r[2], value, 9)
            panel, _ = render_panel.render(art, infos[art])
            buf = io.BytesIO()
            panel.save(buf, "PNG")
            buf.seek(0)
            cap = doc.add_paragraph()
            run = cap.add_run("Vorschau (Stand beim Erstellen):")
            run.font.size, run.font.color.rgb = Pt(8), RGBColor(0x80, 0x80, 0x80)
            doc.add_picture(buf, width=Cm(7.0))
    doc.save(out)
    print(f"{out} ({len(arts)} Produkte)")


def import_(src):
    from docx import Document
    doc = Document(src)
    names = {field for _, field, _ in FIELDS if field not in SLOTS} | set(slot_keys())
    done = 0
    for table in doc.tables:
        cells = [[c.text.strip() for c in r.cells] for r in table.rows]
        if len(cells) < 2 or len(cells[0]) < 3 or cells[1][1] != "Artikelnummer":
            continue
        art = cells[1][2]
        path = PRODUCTS / art / INFO_FILE
        if not path.parent.is_dir():
            print(f"{art}: kein Produktordner, übersprungen")
            continue
        row = {}
        for r in cells[2:]:
            field, value = r[1], " ".join(r[2].split())
            if field not in names:
                print(f"{art}: unbekanntes Feld {field!r} übersprungen")
                continue
            row["Neu (ja/nein)" if field == "Neu" else field] = value
        old = read_info_file(path) if path.exists() else {}
        changed = sorted(k for k in set(old) | set(row) if old.get(k, "") != row.get(k, ""))
        if changed:
            write_info_file(path, art, row)
            print(f"{art}: geändert -- {', '.join(changed)}")
        done += 1
    print(f"{done} Produkte gelesen")


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export")
    e.add_argument("--out", default=str(PRODUCTS / "Produktinfos.docx"))
    i = sub.add_parser("import")
    i.add_argument("docx")
    args = parser.parse_args()
    if args.cmd == "export":
        export(args.out)
    else:
        import_(args.docx)


if __name__ == "__main__":
    main()
